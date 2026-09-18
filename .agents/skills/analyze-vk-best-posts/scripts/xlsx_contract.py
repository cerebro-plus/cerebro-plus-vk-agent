from __future__ import annotations

import math
import os
import tempfile
import zipfile
import datetime as dt
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from common import InputValidationError


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"x": MAIN_NS, "r": REL_NS, "p": PACKAGE_REL_NS}


def _sheet_targets(archive: zipfile.ZipFile) -> dict[str, str]:
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {
        item.attrib["Id"]: item.attrib["Target"].lstrip("/")
        for item in relationships.findall("p:Relationship", NS)
    }
    result: dict[str, str] = {}
    for sheet in workbook.findall("x:sheets/x:sheet", NS):
        relation_id = sheet.attrib[f"{{{REL_NS}}}id"]
        target = targets[relation_id]
        if not target.startswith("xl/"):
            target = "xl/" + target
        result[str(sheet.attrib["name"])] = target.replace("\\", "/")
    return result


def finalize_freeze_panes(
    path: Path,
    freezes: dict[str, tuple[int, str]],
) -> None:
    if not path.is_file():
        raise InputValidationError("XLSX output is missing")
    with zipfile.ZipFile(path, "r") as source:
        targets = _sheet_targets(source)
        replacements: dict[str, bytes] = {}
        for sheet_name, (rows, top_left) in freezes.items():
            target = targets.get(sheet_name)
            if not target:
                raise InputValidationError(f"Missing worksheet for freeze panes: {sheet_name}")
            root = ET.fromstring(source.read(target))
            sheet_views = root.find(f"{{{MAIN_NS}}}sheetViews")
            if sheet_views is None:
                sheet_views = ET.Element(f"{{{MAIN_NS}}}sheetViews")
                root.insert(0, sheet_views)
            sheet_view = sheet_views.find(f"{{{MAIN_NS}}}sheetView")
            if sheet_view is None:
                sheet_view = ET.SubElement(sheet_views, f"{{{MAIN_NS}}}sheetView")
                sheet_view.set("workbookViewId", "0")
            for pane in list(sheet_view.findall(f"{{{MAIN_NS}}}pane")):
                sheet_view.remove(pane)
            pane = ET.Element(
                f"{{{MAIN_NS}}}pane",
                {
                    "ySplit": str(rows),
                    "topLeftCell": top_left,
                    "activePane": "bottomLeft",
                    "state": "frozen",
                },
            )
            sheet_view.insert(0, pane)
            replacements[target] = ET.tostring(
                root, encoding="utf-8", xml_declaration=True
            )
        entries = [(item, source.read(item.filename)) for item in source.infolist()]
    handle, temp_name = tempfile.mkstemp(
        prefix=path.stem + "-", suffix=".xlsx", dir=path.parent
    )
    os.close(handle)
    temp_path = Path(temp_name)
    try:
        with zipfile.ZipFile(temp_path, "w") as destination:
            for item, data in entries:
                destination.writestr(item, replacements.get(item.filename, data))
        with zipfile.ZipFile(temp_path, "r") as check:
            if check.testzip() is not None:
                raise InputValidationError("XLSX ZIP integrity check failed")
        os.replace(temp_path, path)
    finally:
        temp_path.unlink(missing_ok=True)


def _same_number(actual: Any, expected: Any) -> bool:
    if actual is None or expected is None:
        return actual is None and expected is None
    try:
        return math.isclose(float(actual), float(expected), abs_tol=1e-12)
    except (TypeError, ValueError):
        return False


def _formula_has_url(value: Any, url: str) -> bool:
    return isinstance(value, str) and value.startswith("=HYPERLINK(") and url in value


def inspect_xlsx(
    path: Path,
    dataset: dict[str, Any],
    expected_sheets: list[str],
    best_columns: list[str],
    checked_columns: list[str],
    unavailable_columns: list[str],
) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise InputValidationError("openpyxl is required for XLSX validation") from exc
    try:
        with zipfile.ZipFile(path, "r") as archive:
            if archive.testzip() is not None:
                raise InputValidationError("XLSX ZIP integrity check failed")
            embedded_images = len(
                [name for name in archive.namelist() if name.startswith("xl/media/")]
            )
        workbook = load_workbook(path, data_only=False, read_only=False)
    except Exception as exc:
        if isinstance(exc, InputValidationError):
            raise
        raise InputValidationError(f"Cannot reopen XLSX: {type(exc).__name__}") from exc

    issues: list[str] = []
    if workbook.sheetnames != expected_sheets:
        issues.append("sheet_order")
    if issues:
        raise InputValidationError("XLSX contract mismatch: " + ",".join(issues))

    summary = workbook["Сводка"]
    best = workbook["Лучшие посты"]
    checked = workbook["Проверенные посты"]
    missing = workbook["Недоступно"]
    if [best.cell(4, col).value for col in range(1, 15)] != best_columns:
        issues.append("best_columns")
    if [checked.cell(4, col).value for col in range(1, 18)] != checked_columns:
        issues.append("checked_columns")
    if [missing.cell(4, col).value for col in range(1, 4)] != unavailable_columns:
        issues.append("unavailable_columns")

    type_order = {"Прямой": 0, "Косвенный": 1, "За внимание": 2}
    expected_best = sorted(
        [item for item in dataset["posts"] if item.get("is_best")],
        key=lambda item: (
            type_order.get(item["competitor_type"], 9),
            -float(item.get("er") if item.get("er") is not None else -1),
        ),
    )
    expected_checked = sorted(
        dataset["posts"],
        key=lambda item: (
            type_order.get(item["competitor_type"], 9),
            -int(item["timestamp"]),
        ),
    )

    for index, post in enumerate(expected_best, 5):
        values = [best.cell(index, col).value for col in range(1, 15)]
        if values[0] != post["community_name"] or values[1] != post["competitor_type"]:
            issues.append(f"best_identity:{post['post_id']}")
        if not _formula_has_url(values[2], post["url"]):
            issues.append(f"best_link:{post['post_id']}")
        if (values[3] or "") != post["text"] or values[4] != post["post_type"]:
            issues.append(f"best_content:{post['post_id']}")
        if values[7] != post["likes"] or values[8] != post["comments"]:
            issues.append(f"best_counts:{post['post_id']}")
        if not _same_number(values[9], post["er"]) or not _same_number(
            values[10], post["benchmark"]
        ):
            issues.append(f"best_metrics:{post['post_id']}")
        expected_transcript = (
            post["media"].get("transcript") or ""
            if post["post_type"] == "Клип"
            else ""
        )
        if (values[11] or "") != expected_transcript:
            issues.append(f"best_transcript:{post['post_id']}")
        if str(values[13]) != str(post["post_id"]):
            issues.append(f"best_post_id:{post['post_id']}")
        if best.cell(index, 10).number_format != "0.00%" or best.cell(
            index, 11
        ).number_format != "0.00%":
            issues.append(f"best_percent_format:{post['post_id']}")

    for index, post in enumerate(expected_checked, 5):
        values = [checked.cell(index, col).value for col in range(1, 18)]
        if values[0] != post["community_name"] or values[2] != post["followers"]:
            issues.append(f"checked_identity:{post['post_id']}")
        if not _formula_has_url(values[1], post["community_vk"]) or not _formula_has_url(
            values[4], post["url"]
        ):
            issues.append(f"checked_link:{post['post_id']}")
        expected_date = dt.datetime.fromisoformat(post["date"].replace("Z", "+00:00"))
        actual_date = values[3]
        if not isinstance(actual_date, dt.datetime) or actual_date.replace(
            tzinfo=dt.timezone.utc
        ) != expected_date.astimezone(dt.timezone.utc):
            issues.append(f"checked_date:{post['post_id']}")
        expected_values = [
            post["text"],
            post["post_type"],
            "Да" if post["poll"] else "Нет",
            "Да" if post["giveaway"] else "Нет",
            post["likes"],
            post["comments"],
            post["reposts"],
            post["views"],
        ]
        actual_values = values[5:13]
        actual_values[0] = actual_values[0] or ""
        if actual_values != expected_values:
            issues.append(f"checked_content:{post['post_id']}")
        if not _same_number(values[13], post["er"]) or not _same_number(
            values[14], post["benchmark"]
        ):
            issues.append(f"checked_metrics:{post['post_id']}")
        if values[15] != post["result"] or values[16] != post["data_status"]:
            issues.append(f"checked_status:{post['post_id']}")
        if checked.cell(index, 14).number_format != "0.00%" or checked.cell(
            index, 15
        ).number_format != "0.00%":
            issues.append(f"checked_percent_format:{post['post_id']}")

    unavailable = dataset.get("unavailable") or []
    if unavailable:
        for index, item in enumerate(unavailable, 5):
            if missing.cell(index, 1).value != item["object"]:
                issues.append(f"unavailable_object:{index}")
            if not _formula_has_url(missing.cell(index, 2).value, item.get("link") or ""):
                issues.append(f"unavailable_link:{index}")
            if missing.cell(index, 3).value != item["reason"]:
                issues.append(f"unavailable_reason:{index}")
    elif missing.cell(5, 1).value != "Нет":
        issues.append("unavailable_empty_marker")
    expected_rows = {
        "Лучшие посты": max(4, 4 + len(expected_best)),
        "Проверенные посты": max(4, 4 + len(expected_checked)),
        "Недоступно": 4 + max(1, len(unavailable)),
    }
    for sheet_name, expected_row in expected_rows.items():
        if workbook[sheet_name].max_row != expected_row:
            issues.append(f"row_count:{sheet_name}")

    expected_freezes = {
        "Сводка": "A3",
        "Лучшие посты": "A5",
        "Проверенные посты": "A5",
        "Недоступно": "A5",
    }
    for sheet_name, freeze in expected_freezes.items():
        if str(workbook[sheet_name].freeze_panes) != freeze:
            issues.append(f"freeze:{sheet_name}")
    for sheet in (best, checked, missing):
        if len(sheet.tables) != 1:
            issues.append(f"filter_table:{sheet.title}")
        if not all(
            sheet.cell(4, column).alignment.wrap_text
            for column in range(1, sheet.max_column + 1)
        ):
            issues.append(f"header_wrap:{sheet.title}")
        if sheet.max_row >= 5 and not all(
            sheet.cell(row, column).alignment.wrap_text
            for row in range(5, sheet.max_row + 1)
            for column in range(1, sheet.max_column + 1)
        ):
            issues.append(f"body_wrap:{sheet.title}")

    expected_images = sum(
        1
        for item in expected_best
        if (item.get("media") or {}).get("collage_path")
    )
    if embedded_images != expected_images:
        issues.append("embedded_image_count")
    if issues:
        raise InputValidationError("XLSX contract mismatch: " + ",".join(issues))
    return {
        "passed": True,
        "sheets": workbook.sheetnames,
        "best_rows": len(expected_best),
        "checked_rows": len(expected_checked),
        "unavailable_rows": len(unavailable),
        "embedded_images": embedded_images,
        "freeze_panes": expected_freezes,
        "filters": {
            "Лучшие посты": len(best.tables),
            "Проверенные посты": len(checked.tables),
            "Недоступно": len(missing.tables),
        },
    }
